from flask import Blueprint
from flask import (
    render_template, request, redirect, url_for, session, flash,
    current_app, make_response
)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from datetime import datetime
import os
import json
import csv
import io

from core.db import get_db, migrate_db
from core.helpers import (
    login_required, role_required, get_setting, set_setting,
    format_money, log_activity, get_redirect_for_role
)


admin = Blueprint('admin', __name__)

@admin.route('/admin')
@login_required
@role_required('admin')
def admin_dashboard():
    db = get_db()
    migrate_db()

    # Core counts
    total_users = db.execute('SELECT COUNT(*) FROM users').fetchone()[0]
    total_agents = db.execute("SELECT COUNT(*) FROM users WHERE role = 'agent'").fetchone()[0]
    total_customers = db.execute("SELECT COUNT(*) FROM users WHERE role = 'customer'").fetchone()[0]
    total_properties = db.execute('SELECT COUNT(*) FROM properties').fetchone()[0]
    available_properties = db.execute("SELECT COUNT(*) FROM properties WHERE status = 'available'").fetchone()[0]
    total_clients = db.execute('SELECT COUNT(*) FROM clients').fetchone()[0]
    total_appointments = db.execute('SELECT COUNT(*) FROM appointments').fetchone()[0]

    # Total property value
    total_property_value = db.execute('SELECT COALESCE(SUM(price), 0) FROM properties').fetchone()[0] or 0

    # Total deal value & closed deals
    total_deal_value = db.execute('SELECT COALESCE(SUM(value), 0) FROM deals').fetchone()[0] or 0
    closed_deals = db.execute("SELECT COUNT(*) FROM deals WHERE stage = 'closed'").fetchone()[0]
    closed_deal_value = db.execute("SELECT COALESCE(SUM(value), 0) FROM deals WHERE stage = 'closed'").fetchone()[0] or 0

    # Deals closed this month
    deals_closed_this_month = db.execute(
        """SELECT COUNT(*) FROM deals
           WHERE stage = 'closed'
           AND strftime('%Y-%m', updated_at) = strftime('%Y-%m', 'now')"""
    ).fetchone()[0]

    # Revenue / commission earned (closed deals)
    commission_earned = db.execute(
        """SELECT COALESCE(SUM(value * commission_rate / 100.0), 0)
           FROM deals WHERE stage = 'closed'"""
    ).fetchone()[0] or 0

    # Lead conversion rate
    total_leads_or_clients = db.execute('SELECT COUNT(*) FROM clients').fetchone()[0]
    closed_clients = db.execute("SELECT COUNT(*) FROM clients WHERE status = 'closed'").fetchone()[0]
    lead_conversion_rate = round((closed_clients / total_leads_or_clients * 100), 1) if total_leads_or_clients else 0.0

    # Month-over-month growth (properties added)
    props_this_month = db.execute(
        "SELECT COUNT(*) FROM properties WHERE strftime('%Y-%m', created_at) = strftime('%Y-%m', 'now')"
    ).fetchone()[0]
    props_last_month = db.execute(
        """SELECT COUNT(*) FROM properties
           WHERE strftime('%Y-%m', created_at) = strftime('%Y-%m', date('now', '-1 month'))"""
    ).fetchone()[0]
    if props_last_month > 0:
        mom_growth = round(((props_this_month - props_last_month) / props_last_month) * 100, 1)
    else:
        mom_growth = 100.0 if props_this_month > 0 else 0.0

    deals_this_month = db.execute(
        "SELECT COUNT(*) FROM deals WHERE strftime('%Y-%m', created_at) = strftime('%Y-%m', 'now')"
    ).fetchone()[0]
    deals_last_month = db.execute(
        """SELECT COUNT(*) FROM deals
           WHERE strftime('%Y-%m', created_at) = strftime('%Y-%m', date('now', '-1 month'))"""
    ).fetchone()[0]

    stats = {
        'total_users': total_users,
        'total_agents': total_agents,
        'total_customers': total_customers,
        'total_properties': total_properties,
        'available_properties': available_properties,
        'total_clients': total_clients,
        'total_appointments': total_appointments,
        'total_property_value': total_property_value,
        'total_deal_value': total_deal_value,
        'closed_deals': closed_deals,
        'closed_deal_value': closed_deal_value,
        'deals_closed_this_month': deals_closed_this_month,
        'commission_earned': commission_earned,
        'lead_conversion_rate': lead_conversion_rate,
        'mom_growth': mom_growth,
        'props_this_month': props_this_month,
        'props_last_month': props_last_month,
        'deals_this_month': deals_this_month,
        'deals_last_month': deals_last_month,
    }

    # Pending work
    pending_raw = db.execute(
        """SELECT t.id, t.title, t.due_date, t.priority, t.status,
                  u.full_name as assigned_name
           FROM tasks t
           LEFT JOIN users u ON t.assigned_to = u.id
           WHERE t.status IN ('pending', 'in_progress', 'open')
           ORDER BY
             CASE t.priority WHEN 'high' THEN 1 WHEN 'medium' THEN 2 ELSE 3 END,
             t.due_date ASC
           LIMIT 8"""
    ).fetchall()

    pending_tasks = []
    today = datetime.now().date()
    for i, t in enumerate(pending_raw, 1):
        days_left = None
        days_label = '—'
        if t['due_date']:
            try:
                due = datetime.strptime(str(t['due_date'])[:10], '%Y-%m-%d').date()
                days_left = (due - today).days
                if days_left < 0:
                    days_label = f'{abs(days_left)}d overdue'
                elif days_left == 0:
                    days_label = 'Due today'
                elif days_left == 1:
                    days_label = '1 day left'
                else:
                    days_label = f'{days_left} days left'
            except Exception:
                days_label = str(t['due_date'])
        pending_tasks.append({
            'id': t['id'],
            'sl': i,
            'title': t['title'],
            'due_date': t['due_date'],
            'priority': t['priority'] or 'medium',
            'status': t['status'],
            'assigned_name': t['assigned_name'],
            'days_label': days_label,
            'days_left': days_left,
        })

    # Recent activity feed
    activity_raw = db.execute(
        """SELECT a.id, a.action, a.entity_type, a.entity_id, a.details, a.created_at,
                  u.full_name as user_name
           FROM activity_log a
           LEFT JOIN users u ON a.user_id = u.id
           ORDER BY a.created_at DESC
           LIMIT 12"""
    ).fetchall()

    def time_ago(ts):
        if not ts:
            return ''
        try:
            if isinstance(ts, str):
                for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%dT%H:%M:%S', '%Y-%m-%d %H:%M:%S.%f'):
                    try:
                        dt = datetime.strptime(ts[:26].replace('T', ' '), fmt if '%f' not in fmt else '%Y-%m-%d %H:%M:%S.%f')
                        break
                    except ValueError:
                        dt = None
                if dt is None:
                    return str(ts)
            else:
                dt = ts
            diff = datetime.now() - dt
            secs = int(diff.total_seconds())
            if secs < 60:
                return 'just now'
            if secs < 3600:
                m = secs // 60
                return f'{m} min ago' if m == 1 else f'{m} mins ago'
            if secs < 86400:
                h = secs // 3600
                return f'{h} hour ago' if h == 1 else f'{h} hours ago'
            d = secs // 86400
            return f'{d} day ago' if d == 1 else f'{d} days ago'
        except Exception:
            return str(ts)

    activity_feed = []
    for a in activity_raw:
        action = a['action'] or ''
        details = a['details'] or ''
        if details and details not in action:
            change = f'{action} {details}'.strip()
        else:
            change = action
        activity_feed.append({
            'user_name': a['user_name'] or 'System',
            'change': change,
            'time_ago': time_ago(a['created_at']),
            'entity_type': a['entity_type'],
        })

    currency = 'INR'

    current_user = db.execute(
        'SELECT id, full_name, email, role, phone, avatar FROM users WHERE id = ?',
        (session.get('user_id'),)
    ).fetchone()

    return render_template(
        'admin.html',
        stats=stats,
        pending_tasks=pending_tasks,
        activity_feed=activity_feed,
        currency=currency,
        format_money=format_money,
        current_user=current_user,
    )


@admin.route('/admin/users')
@login_required
@role_required('admin')
def admin_users():
    db = get_db()
    migrate_db()
    q = request.args.get('q', '').strip()
    role_filter = request.args.get('role', '').strip()
    status_filter = request.args.get('status', '').strip()
    # Online = last_seen within 10 minutes
    online_sql = (
        "last_seen IS NOT NULL AND "
        "datetime(last_seen) >= datetime('now', '-10 minutes')"
    )
    sql = (
        'SELECT id, full_name, email, role, phone, created_at, is_active, last_seen, avatar, '
        f'CASE WHEN {online_sql} THEN 1 ELSE 0 END AS is_online '
        'FROM users WHERE 1=1'
    )
    params = []
    if q:
        sql += ' AND (full_name LIKE ? OR email LIKE ? OR phone LIKE ?)'
        params.extend([f'%{q}%', f'%{q}%', f'%{q}%'])
    if role_filter in ('admin', 'agent', 'customer'):
        sql += ' AND role = ?'
        params.append(role_filter)
    if status_filter == 'active':
        sql += ' AND is_active = 1'
    elif status_filter == 'inactive':
        sql += ' AND is_active = 0'
    elif status_filter == 'online':
        sql += f' AND ({online_sql})'
    elif status_filter == 'offline':
        sql += f' AND NOT ({online_sql})'
    sql += ' ORDER BY created_at ASC'
    users = db.execute(sql, params).fetchall()
    total = db.execute('SELECT COUNT(*) FROM users').fetchone()[0]
    active_count = db.execute('SELECT COUNT(*) FROM users WHERE is_active = 1').fetchone()[0]
    online_count = db.execute(
        f'SELECT COUNT(*) FROM users WHERE {online_sql}'
    ).fetchone()[0]
    pct = round((active_count / total) * 100) if total else 0
    stats = {
        'total': total,
        'admins': db.execute("SELECT COUNT(*) FROM users WHERE role = 'admin'").fetchone()[0],
        'agents': db.execute("SELECT COUNT(*) FROM users WHERE role = 'agent'").fetchone()[0],
        'customers': db.execute("SELECT COUNT(*) FROM users WHERE role = 'customer'").fetchone()[0],
        'active': active_count,
        'online': online_count,
        'active_pct': pct,
    }
    view = request.args.get('view', '').strip()
    return render_template(
        'admin_users.html', users=users, q=q, role_filter=role_filter,
        status_filter=status_filter, stats=stats, view=view
    )


@admin.route('/admin/users/bulk', methods=['POST'])
@login_required
@role_required('admin')
def admin_bulk_users():
    db = get_db()
    action = request.form.get('action', '').strip()
    ids = request.form.getlist('user_ids')
    ids = [int(i) for i in ids if str(i).isdigit()]
    if not ids:
        flash('No users selected.', 'warning')
        return redirect(url_for('admin.admin_users'))
    self_id = session.get('user_id')
    ids = [i for i in ids if i != self_id]
    if not ids:
        flash('You cannot apply bulk actions to your own account.', 'warning')
        return redirect(url_for('admin.admin_users'))
    placeholders = ','.join('?' * len(ids))
    if action == 'delete':
        db.execute(f'DELETE FROM users WHERE id IN ({placeholders})', ids)
        db.commit()
        flash(f'{len(ids)} user(s) deleted.', 'success')
    else:
        flash('Unknown bulk action.', 'danger')
    return redirect(url_for('admin.admin_users'))


@admin.route('/presence/offline', methods=['POST'])
def presence_offline():
    """Mark current user offline (called on tab/app close via sendBeacon)."""
    uid = session.get('user_id')
    if uid:
        try:
            db = get_db()
            db.execute("UPDATE users SET last_seen = NULL WHERE id = ?", (uid,))
            db.commit()
        except Exception:
            pass
    return ('', 204)


USER_IMPORT_FORMAT = """
============================================================
 ESTACLE USER IMPORT FORMAT (CSV / Excel)
============================================================
 Required columns (header row must match exactly):

   full_name,email,role,phone,password

 Rules:
   - role must be one of: admin | agent | customer
   - password min 6 characters
   - email must be unique
   - phone is optional (can be empty)

 Example CSV:
------------------------------------------------------------
full_name,email,role,phone,password
John Agent,john.agent@estacle.com,agent,+1-555-0101,agent123
Sarah Customer,sarah@estacle.com,customer,+1-555-0102,customer123
Admin Two,admin2@estacle.com,admin,,admin12345
------------------------------------------------------------
 Excel: same headers in row 1, data from row 2. Sheet 1 only.
============================================================
"""


@admin.route('/admin/users/add', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def admin_add_user():
    if request.method == 'POST':
        full_name = request.form.get('full_name', '').strip()
        email = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        role = request.form.get('role', 'agent')
        phone = request.form.get('phone', '').strip()
        alt_phone = request.form.get('alt_phone', '').strip()
        job_title = request.form.get('job_title', '').strip()
        department = request.form.get('department', '').strip()
        employee_id = request.form.get('employee_id', '').strip()
        license_no = request.form.get('license_no', '').strip()
        commission_raw = request.form.get('commission_rate', '').strip()
        address = request.form.get('address', '').strip()
        city = request.form.get('city', '').strip()
        state = request.form.get('state', '').strip()
        country = request.form.get('country', '').strip()
        postal_code = request.form.get('postal_code', '').strip()
        notes = request.form.get('notes', '').strip()
        hire_date = request.form.get('hire_date', '').strip()
        timezone = request.form.get('timezone', '').strip()
        is_active = 0 if request.form.get('is_active') == '0' else 1

        commission_rate = None
        if commission_raw:
            try:
                commission_rate = float(commission_raw)
            except ValueError:
                flash('Commission rate must be a number.', 'danger')
                return redirect(url_for('admin.admin_add_user'))

        if not full_name or not email or not password or role not in ('admin', 'agent', 'customer'):
            flash('Please fill all required fields correctly.', 'danger')
            return redirect(url_for('admin.admin_add_user'))
        if len(password) < 6:
            flash('Password must be at least 6 characters.', 'danger')
            return redirect(url_for('admin.admin_add_user'))

        db = get_db()
        migrate_db()
        existing = db.execute('SELECT id FROM users WHERE email = ?', (email,)).fetchone()
        if existing:
            flash('Email already exists.', 'danger')
            return redirect(url_for('admin.admin_add_user'))

        cur = db.execute(
            "INSERT INTO users ("
            "email, password, full_name, role, phone, is_active, "
            "alt_phone, job_title, department, employee_id, license_no, "
            "commission_rate, address, city, state, country, postal_code, "
            "notes, hire_date, timezone"
            ") VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                email, generate_password_hash(password), full_name, role, phone, is_active,
                alt_phone or None, job_title or None, department or None, employee_id or None,
                license_no or None, commission_rate, address or None, city or None,
                state or None, country or None, postal_code or None, notes or None,
                hire_date or None, timezone or None
            )
        )
        db.commit()
        new_id = cur.lastrowid

        avatar_file = request.files.get('avatar')
        if avatar_file and avatar_file.filename:
            av_dir = os.path.join(current_app.root_path, 'static', 'uploads', 'avatars')
            os.makedirs(av_dir, exist_ok=True)
            ext = os.path.splitext(secure_filename(avatar_file.filename))[1].lower()
            if ext in ('.jpg', '.jpeg', '.png', '.gif', '.webp'):
                avatar_name = 'user_{}_{}{}'.format(new_id, int(datetime.now().timestamp()), ext)
                avatar_file.save(os.path.join(av_dir, avatar_name))
                db.execute('UPDATE users SET avatar = ? WHERE id = ?', (avatar_name, new_id))
                db.commit()

        log_activity('create_user', 'user', new_id, 'Created user {} ({})'.format(full_name, role))
        flash('User {} created successfully.'.format(full_name), 'success')
        return redirect(url_for('admin.admin_users'))

    print(USER_IMPORT_FORMAT)
    return render_template('admin_add_user.html', import_format=USER_IMPORT_FORMAT)


@admin.route('/admin/users/import', methods=['POST'])
@login_required
@role_required('admin')
def admin_import_users():
    print(USER_IMPORT_FORMAT)
    f = request.files.get('import_file')
    if not f or not f.filename:
        flash('Please choose a CSV or Excel file.', 'danger')
        return redirect(url_for('admin.admin_add_user'))

    filename = secure_filename(f.filename).lower()
    rows = []
    try:
        if filename.endswith('.csv'):
            raw = f.read().decode('utf-8-sig')
            reader = csv.DictReader(io.StringIO(raw))
            rows = list(reader)
        elif filename.endswith('.xlsx') or filename.endswith('.xls'):
            try:
                import openpyxl
            except ImportError:
                flash('Excel import requires openpyxl. Run: pip install openpyxl', 'danger')
                return redirect(url_for('admin.admin_add_user'))
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            ws = wb.active
            data = list(ws.iter_rows(values_only=True))
            if not data:
                flash('Excel file is empty.', 'danger')
                return redirect(url_for('admin.admin_add_user'))
            headers = [str(h).strip().lower() if h is not None else '' for h in data[0]]
            for row in data[1:]:
                if row is None or all(c is None or str(c).strip() == '' for c in row):
                    continue
                item = {}
                for i, h in enumerate(headers):
                    if i < len(row):
                        item[h] = '' if row[i] is None else str(row[i]).strip()
                rows.append(item)
        else:
            flash('Unsupported file type. Use .csv or .xlsx', 'danger')
            return redirect(url_for('admin.admin_add_user'))
    except Exception as e:
        flash('Could not read file: {}'.format(str(e)), 'danger')
        return redirect(url_for('admin.admin_add_user'))

    if not rows:
        flash('No data rows found in file.', 'warning')
        return redirect(url_for('admin.admin_add_user'))

    db = get_db()
    created = 0
    skipped = 0
    errors = []
    required_keys = ('full_name', 'email', 'role', 'password')

    for idx, row in enumerate(rows, start=2):
        # normalize keys
        row = { (k or '').strip().lower(): (v or '').strip() if isinstance(v, str) else ('' if v is None else str(v).strip()) for k, v in row.items() }
        full_name = row.get('full_name', '').strip()
        email = row.get('email', '').strip().lower()
        role = row.get('role', '').strip().lower()
        phone = row.get('phone', '').strip()
        password = row.get('password', '').strip()

        if not full_name and not email:
            continue
        if not full_name or not email or not password or role not in ('admin', 'agent', 'customer'):
            skipped += 1
            errors.append('Row {}: missing/invalid fields'.format(idx))
            continue
        if len(password) < 6:
            skipped += 1
            errors.append('Row {}: password too short'.format(idx))
            continue
        existing = db.execute('SELECT id FROM users WHERE email = ?', (email,)).fetchone()
        if existing:
            skipped += 1
            errors.append('Row {}: email already exists ({})'.format(idx, email))
            continue
        try:
            db.execute(
                'INSERT INTO users (email, password, full_name, role, phone) VALUES (?, ?, ?, ?, ?)',
                (email, generate_password_hash(password), full_name, role, phone)
            )
            created += 1
        except Exception as e:
            skipped += 1
            errors.append('Row {}: {}'.format(idx, str(e)))

    db.commit()
    if created:
        log_activity('import_users', 'user', None, 'Imported {} users'.format(created))
    msg = 'Import finished: {} created, {} skipped.'.format(created, skipped)
    if errors:
        msg += ' Issues: ' + '; '.join(errors[:5])
        if len(errors) > 5:
            msg += ' (+{} more)'.format(len(errors) - 5)
    flash(msg, 'success' if created else 'warning')
    print('[IMPORT] ' + msg)
    return redirect(url_for('admin.admin_users'))


@admin.route('/admin/users/import/template.csv')
@login_required
@role_required('admin')
def admin_users_import_template():
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['full_name', 'email', 'role', 'phone', 'password'])
    writer.writerow(['John Agent', 'john.agent@estacle.com', 'agent', '+1-555-0101', 'agent123'])
    writer.writerow(['Sarah Customer', 'sarah@estacle.com', 'customer', '+1-555-0102', 'customer123'])
    writer.writerow(['Admin Two', 'admin2@estacle.com', 'admin', '', 'admin12345'])
    resp = make_response(output.getvalue())
    resp.headers['Content-Type'] = 'text/csv; charset=utf-8'
    resp.headers['Content-Disposition'] = 'attachment; filename=estacle_users_import_template.csv'
    return resp



@admin.route('/admin/users/toggle/<int:user_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_toggle_user(user_id):
    db = get_db()
    user = db.execute('SELECT id, is_active, role FROM users WHERE id = ?', (user_id,)).fetchone()
    if not user:
        flash('User not found.', 'danger')
        return redirect(url_for('admin.admin_users'))
    if user['id'] == session['user_id']:
        flash('You cannot deactivate yourself.', 'warning')
        return redirect(url_for('admin.admin_users'))

    new_status = 0 if user['is_active'] else 1
    db.execute('UPDATE users SET is_active = ? WHERE id = ?', (new_status, user_id))
    db.commit()
    flash('User status updated.', 'success')
    return redirect(url_for('admin.admin_user_detail', user_id=user_id))


@admin.route('/admin/users/<int:user_id>')
@login_required
@role_required('admin')
def admin_user_detail(user_id):
    db = get_db()
    migrate_db()
    online_sql = (
        "last_seen IS NOT NULL AND "
        "datetime(last_seen) >= datetime('now', '-10 minutes')"
    )
    user = db.execute(
        'SELECT *, CASE WHEN ' + online_sql + ' THEN 1 ELSE 0 END AS is_online '
        'FROM users WHERE id = ?',
        (user_id,)
    ).fetchone()
    if not user:
        flash('User not found.', 'danger')
        return redirect(url_for('admin.admin_users'))

    activity = db.execute(
        'SELECT action, entity_type, entity_id, details, created_at '
        'FROM activity_log WHERE user_id = ? ORDER BY created_at DESC LIMIT 20',
        (user_id,)
    ).fetchall()

    return render_template(
        'admin_user_detail.html',
        user=user,
        activity=activity,
    )


@admin.route('/admin/users/<int:user_id>/delete', methods=['POST'])
@login_required
@role_required('admin')
def admin_delete_user(user_id):
    db = get_db()
    user = db.execute('SELECT id, full_name FROM users WHERE id = ?', (user_id,)).fetchone()
    if not user:
        flash('User not found.', 'danger')
        return redirect(url_for('admin.admin_users'))
    if user['id'] == session.get('user_id'):
        flash('You cannot delete your own account.', 'warning')
        return redirect(url_for('admin.admin_user_detail', user_id=user_id))
    db.execute('DELETE FROM users WHERE id = ?', (user_id,))
    db.commit()
    log_activity('delete_user', 'user', user_id, 'Deleted user ' + user['full_name'])
    flash('User ' + user['full_name'] + ' deleted.', 'success')
    return redirect(url_for('admin.admin_users'))


@admin.route('/admin/properties')
@login_required
@role_required('admin')
def admin_properties():
    db = get_db()
    migrate_db()
    q = request.args.get('q', '').strip()
    status_filter = request.args.get('status', '').strip()
    listing_type = request.args.get('listing_type', '').strip()
    prop_type = request.args.get('property_type', '').strip()
    city = request.args.get('city', '').strip()
    agent_id = request.args.get('agent_id', '').strip()
    beds = request.args.get('beds', '').strip()
    pmin = request.args.get('pmin', '').strip()
    pmax = request.args.get('pmax', '').strip()
    date_from = request.args.get('date_from', '').strip()
    date_to = request.args.get('date_to', '').strip()
    view = request.args.get('view', '').strip()

    sql = (
        'SELECT p.*, u.full_name as agent_name '
        'FROM properties p LEFT JOIN users u ON p.agent_id = u.id WHERE 1=1'
    )
    params = []
    if q:
        sql += ' AND (p.title LIKE ? OR p.city LIKE ? OR p.address LIKE ? OR p.description LIKE ?)'
        params.extend(['%' + q + '%'] * 4)
    if status_filter in ('available', 'under_offer', 'sold', 'rented'):
        sql += ' AND p.status = ?'
        params.append(status_filter)
    if listing_type in ('sale', 'rent'):
        sql += ' AND p.listing_type = ?'
        params.append(listing_type)
    if prop_type:
        sql += ' AND p.property_type = ?'
        params.append(prop_type)
    if city:
        sql += ' AND p.city LIKE ?'
        params.append('%' + city + '%')
    if agent_id.isdigit():
        sql += ' AND p.agent_id = ?'
        params.append(int(agent_id))
    if beds.isdigit():
        sql += ' AND p.bedrooms >= ?'
        params.append(int(beds))
    if pmin:
        try:
            sql += ' AND p.price >= ?'
            params.append(float(pmin))
        except ValueError:
            pass
    if pmax:
        try:
            sql += ' AND p.price <= ?'
            params.append(float(pmax))
        except ValueError:
            pass
    def _parse_date(val):
        """Accept YYYY-MM-DD or DD/MM/YYYY."""
        if not val:
            return None
        val = val.strip()
        if len(val) == 10 and val[4] == '-' and val[7] == '-':
            return val
        if len(val) == 10 and val[2] == '/' and val[5] == '/':
            d, m, y = val.split('/')
            return '{}-{}-{}'.format(y, m, d)
        return None

    df = _parse_date(date_from)
    dt = _parse_date(date_to)
    if df:
        sql += ' AND date(p.created_at) >= date(?)'
        params.append(df)
    if dt:
        sql += ' AND date(p.created_at) <= date(?)'
        params.append(dt)
    sql += ' ORDER BY p.created_at ASC'
    properties = db.execute(sql, params).fetchall()

    # Relative "added" labels for table
    now = datetime.now().date()
    props_out = []
    for row in properties:
        d = dict(row)
        created = d.get('created_at')
        label = '—'
        if created:
            try:
                cdt = datetime.strptime(str(created)[:10], '%Y-%m-%d').date()
                days = (now - cdt).days
                if days <= 0:
                    label = 'Today'
                elif days == 1:
                    label = '1 day ago'
                else:
                    label = '{} days ago'.format(days)
            except Exception:
                label = str(created)[:10]
        d['added_label'] = label
        props_out.append(d)
    properties = props_out

    total = db.execute('SELECT COUNT(*) FROM properties').fetchone()[0]
    stats = {
        'total': total,
        'available': db.execute("SELECT COUNT(*) FROM properties WHERE status = 'available'").fetchone()[0],
        'under_offer': db.execute("SELECT COUNT(*) FROM properties WHERE status = 'under_offer'").fetchone()[0],
        'sold': db.execute("SELECT COUNT(*) FROM properties WHERE status = 'sold'").fetchone()[0],
        'rented': db.execute("SELECT COUNT(*) FROM properties WHERE status = 'rented'").fetchone()[0],
        'sale': db.execute("SELECT COUNT(*) FROM properties WHERE listing_type = 'sale'").fetchone()[0],
        'rent': db.execute("SELECT COUNT(*) FROM properties WHERE listing_type = 'rent'").fetchone()[0],
    }
    agents = db.execute(
        "SELECT id, full_name FROM users WHERE role = 'agent' AND is_active = 1 ORDER BY full_name"
    ).fetchall()
    cities = [
        r[0] for r in db.execute(
            "SELECT DISTINCT city FROM properties WHERE city IS NOT NULL AND city != '' ORDER BY city"
        ).fetchall()
    ]

    active_filters = 0
    for v in (q, status_filter, listing_type, prop_type, city, agent_id, beds, pmin, pmax, date_from, date_to):
        if v:
            active_filters += 1

    currency = get_setting('currency', 'INR')
    return render_template(
        'admin_properties.html',
        properties=properties,
        stats=stats,
        agents=agents,
        cities=cities,
        q=q,
        status_filter=status_filter,
        listing_type=listing_type,
        prop_type=prop_type,
        city=city,
        agent_id=agent_id,
        beds=beds,
        pmin=pmin,
        pmax=pmax,
        date_from=date_from,
        date_to=date_to,
        view=view,
        active_filters=active_filters,
        currency=currency,
        format_money=format_money,
    )


@admin.route('/admin/clients')
@login_required
@role_required('admin')
def admin_clients():
    db = get_db()
    clients = db.execute('''
        SELECT c.*, u.full_name as agent_name
        FROM clients c
        LEFT JOIN users u ON c.agent_id = u.id
        ORDER BY c.created_at DESC
    ''').fetchall()
    return render_template('admin_clients.html', clients=clients)


@admin.route('/admin/appointments')
@login_required
@role_required('admin')
def admin_appointments():
    db = get_db()
    appointments = db.execute("""
        SELECT a.*,
               c.full_name as client_name,
               p.title as property_title,
               u.full_name as agent_name
        FROM appointments a
        LEFT JOIN clients c ON a.client_id = c.id
        LEFT JOIN properties p ON a.property_id = p.id
        LEFT JOIN users u ON a.agent_id = u.id
        ORDER BY a.appointment_date DESC
    """).fetchall()
    return render_template('admin_appointments.html', appointments=appointments)


@admin.route('/admin/properties/add', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def admin_add_property():
    db = get_db()
    migrate_db()
    agents = db.execute("SELECT id, full_name FROM users WHERE role = 'agent' AND is_active = 1 ORDER BY full_name").fetchall()
    currency = get_setting('currency', 'INR')

    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        if not title:
            flash('Title is required.', 'danger')
            return redirect(url_for('admin.admin_add_property'))
        try:
            price = float(request.form.get('price') or 0)
            bedrooms = int(request.form.get('bedrooms') or 0) or None
            bathrooms = float(request.form.get('bathrooms') or 0) or None
            area_sqft = float(request.form.get('area_sqft') or 0) or None
            carpet_area = float(request.form.get('carpet_area') or 0) or None
            monthly_rent = float(request.form.get('monthly_rent') or 0) or None
            security_deposit = float(request.form.get('security_deposit') or 0) or None
            maintenance = float(request.form.get('maintenance') or 0) or None
            per_sqft = float(request.form.get('per_sqft') or 0) or None
            agent_id = int(request.form.get('agent_id') or 0) or None
        except ValueError:
            flash('Invalid numeric values.', 'danger')
            return redirect(url_for('admin.admin_add_property'))

        photo_names = []
        files = request.files.getlist('photos')
        for f in files:
            if f and f.filename:
                fname = secure_filename(f.filename)
                unique = '{}_{}'.format(int(__import__('time').time()), fname)
                f.save(os.path.join(current_app.config['UPLOAD_FOLDER'], unique))
                photo_names.append(unique)
        photos_str = ','.join(photo_names) if photo_names else None

        video_names = []
        vfiles = request.files.getlist('videos')
        for f in vfiles:
            if f and f.filename:
                fname = secure_filename(f.filename)
                ext = os.path.splitext(fname)[1].lower()
                if ext not in ('.mp4', '.webm', '.mov', '.ogg', '.m4v'):
                    continue
                unique = 'vid_{}_{}'.format(int(__import__('time').time()), fname)
                f.save(os.path.join(current_app.config['UPLOAD_FOLDER'], unique))
                video_names.append(unique)
        videos_str = ','.join(video_names) if video_names else None

        db.execute(
            """INSERT INTO properties (
                title, description, price, property_type, status, address, city, state, zip_code,
                bedrooms, bathrooms, area_sqft, agent_id, category, listing_type, floor, carpet_area,
                furnishing, parking, amenities, seller_name, seller_phone, monthly_rent, security_deposit,
                maintenance, per_sqft, latitude, longitude, photos, year_built, facing, videos
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                title, request.form.get('description', ''), price,
                request.form.get('property_type', 'house'), request.form.get('status', 'available'),
                request.form.get('address', ''), request.form.get('city', ''), request.form.get('state', ''),
                request.form.get('zip_code', ''), bedrooms, bathrooms, area_sqft, agent_id,
                request.form.get('category', 'residential'), request.form.get('listing_type', 'sale'),
                request.form.get('floor', ''), carpet_area, request.form.get('furnishing', ''),
                request.form.get('parking', ''), request.form.get('amenities', ''),
                request.form.get('seller_name', ''), request.form.get('seller_phone', ''),
                monthly_rent, security_deposit, maintenance, per_sqft,
                request.form.get('latitude', ''), request.form.get('longitude', ''),
                photos_str, request.form.get('year_built', ''), request.form.get('facing', ''), videos_str
            )
        )
        db.commit()
        log_activity('Added property', 'property', None, title)
        flash(f'Property "{title}" published.', 'success')
        return redirect(url_for('admin.admin_properties'))

    print(PROPERTY_IMPORT_FORMAT)
    return render_template('admin_add_property.html', agents=agents, currency=currency)


PROPERTY_IMPORT_FORMAT = """
============================================================
 ESTACLE PROPERTY IMPORT FORMAT (CSV / Excel)
============================================================
 Header row (exact names):

title,description,price,property_type,status,listing_type,category,address,city,state,zip_code,bedrooms,bathrooms,area_sqft,carpet_area,floor,furnishing,parking,amenities,seller_name,seller_phone,monthly_rent,security_deposit,maintenance,per_sqft,year_built,facing,agent_email

 Rules:
  - title and price required
  - property_type: apartment|house|villa|plot|commercial|office|warehouse
  - status: available|under_offer|sold|rented
  - listing_type: sale|rent
  - category: residential|commercial|industrial
  - agent_email: optional, must match an existing agent account
  - numeric fields can be empty

 Example row:
3BHK Sea View,Corner unit with balcony,12500000,apartment,available,sale,residential,Marine Drive,Mumbai,MH,400001,3,2,1450,1200,12,furnished,2 covered,"Pool,Gym",Rajesh,+9198xxxxxx,0,0,5000,8620,2019,East,agent@estacle.com
============================================================
"""


@admin.route('/admin/properties/import', methods=['POST'])
@login_required
@role_required('admin')
def admin_import_properties():
    print(PROPERTY_IMPORT_FORMAT)
    f = request.files.get('import_file')
    if not f or not f.filename:
        flash('Please choose a CSV or Excel file.', 'danger')
        return redirect(url_for('admin.admin_add_property'))

    filename = secure_filename(f.filename).lower()
    rows = []
    try:
        if filename.endswith('.csv'):
            raw = f.read().decode('utf-8-sig')
            reader = csv.DictReader(io.StringIO(raw))
            rows = list(reader)
        elif filename.endswith('.xlsx') or filename.endswith('.xls'):
            try:
                import openpyxl
            except ImportError:
                flash('Excel import requires openpyxl. Run: pip install openpyxl', 'danger')
                return redirect(url_for('admin.admin_add_property'))
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            ws = wb.active
            data = list(ws.iter_rows(values_only=True))
            if not data:
                flash('Excel file is empty.', 'danger')
                return redirect(url_for('admin.admin_add_property'))
            headers = [str(h).strip().lower() if h is not None else '' for h in data[0]]
            for row in data[1:]:
                if row is None or all(c is None or str(c).strip() == '' for c in row):
                    continue
                item = {}
                for i, h in enumerate(headers):
                    if i < len(row):
                        item[h] = '' if row[i] is None else str(row[i]).strip()
                rows.append(item)
        else:
            flash('Unsupported file type. Use .csv or .xlsx', 'danger')
            return redirect(url_for('admin.admin_add_property'))
    except Exception as e:
        flash('Could not read file: {}'.format(str(e)), 'danger')
        return redirect(url_for('admin.admin_add_property'))

    if not rows:
        flash('No data rows found.', 'warning')
        return redirect(url_for('admin.admin_add_property'))

    db = get_db()
    migrate_db()
    created = 0
    skipped = 0
    errors = []

    for idx_row, row in enumerate(rows, start=2):
        row = {(k or '').strip().lower(): ('' if v is None else str(v).strip()) for k, v in row.items()}
        title = row.get('title', '').strip()
        if not title:
            skipped += 1
            continue
        try:
            price = float(row.get('price') or 0)
        except ValueError:
            skipped += 1
            errors.append('Row {}: invalid price'.format(idx_row))
            continue

        def num(key, cast=float):
            raw = row.get(key, '').strip()
            if not raw:
                return None
            try:
                return cast(raw)
            except ValueError:
                return None

        agent_id = None
        agent_email = row.get('agent_email', '').strip().lower()
        if agent_email:
            ag = db.execute(
                "SELECT id FROM users WHERE email = ? AND role = 'agent'", (agent_email,)
            ).fetchone()
            if ag:
                agent_id = ag['id']

        status = row.get('status', 'available').strip().lower() or 'available'
        if status not in ('available', 'under_offer', 'sold', 'rented'):
            status = 'available'
        listing_type = row.get('listing_type', 'sale').strip().lower() or 'sale'
        if listing_type not in ('sale', 'rent'):
            listing_type = 'sale'

        try:
            db.execute(
                "INSERT INTO properties ("
                "title, description, price, property_type, status, address, city, state, zip_code, "
                "bedrooms, bathrooms, area_sqft, agent_id, category, listing_type, floor, carpet_area, "
                "furnishing, parking, amenities, seller_name, seller_phone, monthly_rent, security_deposit, "
                "maintenance, per_sqft, year_built, facing"
                ") VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (
                    title,
                    row.get('description', ''),
                    price,
                    row.get('property_type', 'house') or 'house',
                    status,
                    row.get('address', ''),
                    row.get('city', ''),
                    row.get('state', ''),
                    row.get('zip_code', ''),
                    num('bedrooms', int),
                    num('bathrooms', float),
                    num('area_sqft', float),
                    agent_id,
                    row.get('category', 'residential') or 'residential',
                    listing_type,
                    row.get('floor') or None,
                    num('carpet_area', float),
                    row.get('furnishing') or None,
                    row.get('parking') or None,
                    row.get('amenities') or None,
                    row.get('seller_name') or None,
                    row.get('seller_phone') or None,
                    num('monthly_rent', float),
                    num('security_deposit', float),
                    num('maintenance', float),
                    num('per_sqft', float),
                    row.get('year_built') or None,
                    row.get('facing') or None,
                )
            )
            created += 1
        except Exception as e:
            skipped += 1
            errors.append('Row {}: {}'.format(idx_row, str(e)))

    db.commit()
    if created:
        log_activity('import_properties', 'property', None, 'Imported {} properties'.format(created))
    msg = 'Import finished: {} created, {} skipped.'.format(created, skipped)
    if errors:
        msg += ' Issues: ' + '; '.join(errors[:4])
    flash(msg, 'success' if created else 'warning')
    print('[PROPERTY IMPORT] ' + msg)
    return redirect(url_for('admin.admin_properties'))


@admin.route('/admin/properties/import/template.csv')
@login_required
@role_required('admin')
def admin_properties_import_template():
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        'title', 'description', 'price', 'property_type', 'status', 'listing_type', 'category',
        'address', 'city', 'state', 'zip_code', 'bedrooms', 'bathrooms', 'area_sqft', 'carpet_area',
        'floor', 'furnishing', 'parking', 'amenities', 'seller_name', 'seller_phone',
        'monthly_rent', 'security_deposit', 'maintenance', 'per_sqft', 'year_built', 'facing', 'agent_email'
    ])
    writer.writerow([
        '3BHK Sea View', 'Corner unit with balcony', '12500000', 'apartment', 'available', 'sale',
        'residential', 'Marine Drive', 'Mumbai', 'MH', '400001', '3', '2', '1450', '1200',
        '12', 'furnished', '2 covered', 'Pool; Gym', 'Rajesh', '+919800000000',
        '', '', '5000', '8620', '2019', 'East', ''
    ])
    resp = make_response(output.getvalue())
    resp.headers['Content-Type'] = 'text/csv; charset=utf-8'
    resp.headers['Content-Disposition'] = 'attachment; filename=estacle_properties_import_template.csv'
    return resp



@admin.route('/admin/clients/add', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def admin_add_client():
    db = get_db()
    agents = db.execute("SELECT id, full_name FROM users WHERE role = 'agent' AND is_active = 1 ORDER BY full_name").fetchall()

    if request.method == 'POST':
        full_name = request.form.get('full_name', '').strip()
        email = request.form.get('email', '').strip()
        phone = request.form.get('phone', '').strip()
        status = request.form.get('status', 'lead')
        notes = request.form.get('notes', '').strip()
        agent_id = request.form.get('agent_id') or None

        if not full_name:
            flash('Full name is required.', 'danger')
            return redirect(url_for('admin.admin_add_client'))

        agent_id = int(agent_id) if agent_id else None
        db.execute(
            'INSERT INTO clients (full_name, email, phone, status, notes, agent_id) VALUES (?, ?, ?, ?, ?, ?)',
            (full_name, email, phone, status, notes, agent_id)
        )
        db.commit()
        flash(f'Client "{full_name}" added successfully.', 'success')
        return redirect(url_for('admin.admin_clients'))

    return render_template('admin_add_client.html', agents=agents)


@admin.route('/admin/users/edit/<int:user_id>', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def admin_edit_user(user_id):
    db = get_db()
    migrate_db()
    user = db.execute('SELECT * FROM users WHERE id = ?', (user_id,)).fetchone()
    if not user:
        flash('User not found.', 'danger')
        return redirect(url_for('admin.admin_users'))

    if request.method == 'POST':
        full_name = request.form.get('full_name', '').strip()
        email = request.form.get('email', '').strip().lower()
        role = request.form.get('role', user['role'])
        phone = request.form.get('phone', '').strip()
        password = request.form.get('password', '').strip()
        alt_phone = request.form.get('alt_phone', '').strip()
        job_title = request.form.get('job_title', '').strip()
        department = request.form.get('department', '').strip()
        employee_id = request.form.get('employee_id', '').strip()
        license_no = request.form.get('license_no', '').strip()
        commission_raw = request.form.get('commission_rate', '').strip()
        address = request.form.get('address', '').strip()
        city = request.form.get('city', '').strip()
        state = request.form.get('state', '').strip()
        country = request.form.get('country', '').strip()
        postal_code = request.form.get('postal_code', '').strip()
        notes = request.form.get('notes', '').strip()
        hire_date = request.form.get('hire_date', '').strip()
        timezone = request.form.get('timezone', '').strip()

        commission_rate = None
        if commission_raw:
            try:
                commission_rate = float(commission_raw)
            except ValueError:
                flash('Commission rate must be a number.', 'danger')
                return redirect(url_for('admin.admin_edit_user', user_id=user_id))

        if not full_name or not email or role not in ('admin', 'agent', 'customer'):
            flash('Please fill all required fields correctly.', 'danger')
            return redirect(url_for('admin.admin_edit_user', user_id=user_id))

        existing = db.execute('SELECT id FROM users WHERE email = ? AND id != ?', (email, user_id)).fetchone()
        if existing:
            flash('Email already in use by another user.', 'danger')
            return redirect(url_for('admin.admin_edit_user', user_id=user_id))

        avatar_name = user['avatar'] if 'avatar' in user.keys() else None
        avatar_file = request.files.get('avatar')
        if avatar_file and avatar_file.filename:
            av_dir = os.path.join(current_app.root_path, 'static', 'uploads', 'avatars')
            os.makedirs(av_dir, exist_ok=True)
            ext = os.path.splitext(secure_filename(avatar_file.filename))[1].lower()
            if ext in ('.jpg', '.jpeg', '.png', '.gif', '.webp'):
                avatar_name = 'user_{}_{}{}'.format(user_id, int(datetime.now().timestamp()), ext)
                avatar_file.save(os.path.join(av_dir, avatar_name))

        fields = (
            full_name, email, role, phone, avatar_name,
            alt_phone or None, job_title or None, department or None, employee_id or None,
            license_no or None, commission_rate, address or None, city or None,
            state or None, country or None, postal_code or None, notes or None,
            hire_date or None, timezone or None, user_id
        )
        if password:
            if len(password) < 6:
                flash('Password must be at least 6 characters.', 'danger')
                return redirect(url_for('admin.admin_edit_user', user_id=user_id))
            db.execute(
                "UPDATE users SET full_name=?, email=?, role=?, phone=?, avatar=?, "
                "alt_phone=?, job_title=?, department=?, employee_id=?, license_no=?, "
                "commission_rate=?, address=?, city=?, state=?, country=?, postal_code=?, "
                "notes=?, hire_date=?, timezone=?, password=? WHERE id=?",
                (
                    full_name, email, role, phone, avatar_name,
                    alt_phone or None, job_title or None, department or None, employee_id or None,
                    license_no or None, commission_rate, address or None, city or None,
                    state or None, country or None, postal_code or None, notes or None,
                    hire_date or None, timezone or None, generate_password_hash(password), user_id
                )
            )
        else:
            db.execute(
                "UPDATE users SET full_name=?, email=?, role=?, phone=?, avatar=?, "
                "alt_phone=?, job_title=?, department=?, employee_id=?, license_no=?, "
                "commission_rate=?, address=?, city=?, state=?, country=?, postal_code=?, "
                "notes=?, hire_date=?, timezone=? WHERE id=?",
                fields
            )
        db.commit()
        if user_id == session.get('user_id'):
            session['full_name'] = full_name
            session['email'] = email
            session['role'] = role
        log_activity('update_user', 'user', user_id, 'Updated user {}'.format(full_name))
        flash('User updated successfully.', 'success')
        if request.form.get('next') == 'dashboard':
            return redirect(url_for('admin.admin_dashboard'))
        return redirect(url_for('admin.admin_user_detail', user_id=user_id))

    return render_template('admin_edit_user.html', user=user)


@admin.route('/admin/properties/delete/<int:prop_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_delete_property(prop_id):
    db = get_db()
    prop = db.execute('SELECT id, title FROM properties WHERE id = ?', (prop_id,)).fetchone()
    if not prop:
        flash('Property not found.', 'danger')
        return redirect(url_for('admin.admin_properties'))
    db.execute('DELETE FROM properties WHERE id = ?', (prop_id,))
    db.commit()
    flash(f'Property "{prop["title"]}" deleted.', 'success')
    return redirect(url_for('admin.admin_properties'))


@admin.route('/admin/clients/delete/<int:client_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_delete_client(client_id):
    db = get_db()
    client = db.execute('SELECT id, full_name FROM clients WHERE id = ?', (client_id,)).fetchone()
    if not client:
        flash('Client not found.', 'danger')
        return redirect(url_for('admin.admin_clients'))
    db.execute('DELETE FROM clients WHERE id = ?', (client_id,))
    db.commit()
    flash(f'Client "{client["full_name"]}" deleted.', 'success')
    return redirect(url_for('admin.admin_clients'))


@admin.route('/admin/properties/status/<int:prop_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_update_property_status(prop_id):
    status = request.form.get('status', '')
    if status not in ('available', 'under_offer', 'sold', 'rented'):
        flash('Invalid status.', 'danger')
        return redirect(url_for('admin.admin_properties'))
    db = get_db()
    prop = db.execute('SELECT id, title FROM properties WHERE id = ?', (prop_id,)).fetchone()
    if not prop:
        flash('Property not found.', 'danger')
        return redirect(url_for('admin.admin_properties'))
    db.execute('UPDATE properties SET status = ? WHERE id = ?', (status, prop_id))
    db.commit()
    flash(f'Status of "{prop["title"]}" updated to {status}.', 'success')
    return redirect(url_for('admin.admin_properties'))


@admin.route('/admin/clients/status/<int:client_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_update_client_status(client_id):
    status = request.form.get('status', '')
    if status not in ('lead', 'contacted', 'qualified', 'negotiation', 'closed', 'lost'):
        flash('Invalid status.', 'danger')
        return redirect(url_for('admin.admin_clients'))
    db = get_db()
    client = db.execute('SELECT id, full_name FROM clients WHERE id = ?', (client_id,)).fetchone()
    if not client:
        flash('Client not found.', 'danger')
        return redirect(url_for('admin.admin_clients'))
    db.execute('UPDATE clients SET status = ? WHERE id = ?', (status, client_id))
    db.commit()
    flash(f'Status of "{client["full_name"]}" updated to {status}.', 'success')
    return redirect(url_for('admin.admin_clients'))


@admin.route('/admin/appointments/add', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def admin_add_appointment():
    db = get_db()
    clients = db.execute('SELECT id, full_name FROM clients ORDER BY full_name').fetchall()
    properties = db.execute('SELECT id, title FROM properties ORDER BY title').fetchall()
    agents = db.execute("SELECT id, full_name FROM users WHERE role = 'agent' AND is_active = 1 ORDER BY full_name").fetchall()

    if request.method == 'POST':
        client_id = request.form.get('client_id') or None
        property_id = request.form.get('property_id') or None
        agent_id = request.form.get('agent_id') or None
        appointment_date = request.form.get('appointment_date', '').strip()
        notes = request.form.get('notes', '').strip()
        status = request.form.get('status', 'scheduled')

        if not appointment_date:
            flash('Appointment date is required.', 'danger')
            return redirect(url_for('admin.admin_add_appointment'))

        client_id = int(client_id) if client_id else None
        property_id = int(property_id) if property_id else None
        agent_id = int(agent_id) if agent_id else None

        db.execute(
            'INSERT INTO appointments (client_id, property_id, agent_id, appointment_date, status, notes) VALUES (?, ?, ?, ?, ?, ?)',
            (client_id, property_id, agent_id, appointment_date, status, notes)
        )
        db.commit()
        flash('Appointment scheduled successfully.', 'success')
        return redirect(url_for('admin.admin_appointments'))

    return render_template('admin_add_appointment.html', clients=clients, properties=properties, agents=agents)


@admin.route('/admin/analytics')
@login_required
@role_required('admin')
def admin_analytics():
    from datetime import timedelta
    import csv
    import io
    from flask import make_response

    db = get_db()
    migrate_db()

    from_date = (request.args.get('from') or '').strip()
    to_date = (request.args.get('to') or '').strip()
    export = (request.args.get('export') or '').strip().lower()

    today = datetime.now().date()
    if not to_date:
        to_date = today.isoformat()
    if not from_date:
        from_date = (today - timedelta(days=90)).isoformat()

    params = (from_date, to_date)

    total_properties = db.execute(
        "SELECT COUNT(*) FROM properties WHERE date(created_at) >= date(?) AND date(created_at) <= date(?)",
        params
    ).fetchone()[0]
    available = db.execute(
        "SELECT COUNT(*) FROM properties WHERE status = 'available' AND date(created_at) >= date(?) AND date(created_at) <= date(?)",
        params
    ).fetchone()[0]
    closed_deals = db.execute(
        "SELECT COUNT(*) FROM deals WHERE stage = 'closed' AND date(updated_at) >= date(?) AND date(updated_at) <= date(?)",
        params
    ).fetchone()[0]
    total_leads = db.execute(
        "SELECT COUNT(*) FROM clients WHERE date(created_at) >= date(?) AND date(created_at) <= date(?)",
        params
    ).fetchone()[0]
    try:
        total_leads += db.execute(
            "SELECT COUNT(*) FROM leads WHERE date(created_at) >= date(?) AND date(created_at) <= date(?)",
            params
        ).fetchone()[0]
    except Exception:
        pass

    sold = db.execute(
        "SELECT COUNT(*) FROM properties WHERE status = 'sold' AND date(created_at) >= date(?) AND date(created_at) <= date(?)",
        params
    ).fetchone()[0]
    rented = db.execute(
        "SELECT COUNT(*) FROM properties WHERE status = 'rented' AND date(created_at) >= date(?) AND date(created_at) <= date(?)",
        params
    ).fetchone()[0]
    under_offer = db.execute(
        "SELECT COUNT(*) FROM properties WHERE status = 'under_offer' AND date(created_at) >= date(?) AND date(created_at) <= date(?)",
        params
    ).fetchone()[0]
    total_appointments = db.execute(
        "SELECT COUNT(*) FROM appointments WHERE date(created_at) >= date(?) AND date(created_at) <= date(?)",
        params
    ).fetchone()[0]
    total_agents = db.execute("SELECT COUNT(*) FROM users WHERE role = 'agent'").fetchone()[0]
    open_deals = db.execute(
        "SELECT COUNT(*) FROM deals WHERE stage NOT IN ('closed', 'lost') AND date(created_at) >= date(?) AND date(created_at) <= date(?)",
        params
    ).fetchone()[0]
    pipeline_value = db.execute(
        "SELECT COALESCE(SUM(value), 0) FROM deals WHERE stage NOT IN ('closed', 'lost') AND date(created_at) >= date(?) AND date(created_at) <= date(?)",
        params
    ).fetchone()[0] or 0

    stats = {
        'total_properties': total_properties,
        'available': available,
        'closed_deals': closed_deals,
        'total_leads': total_leads,
        'sold': sold,
        'rented': rented,
        'under_offer': under_offer,
        'total_appointments': total_appointments,
        'total_agents': total_agents,
        'open_deals': open_deals,
        'pipeline_value': pipeline_value,
    }

    by_type = db.execute(
        "SELECT property_type, COUNT(*) as cnt FROM properties WHERE date(created_at) >= date(?) AND date(created_at) <= date(?) GROUP BY property_type ORDER BY cnt DESC",
        params
    ).fetchall()
    by_city = db.execute(
        "SELECT city, COUNT(*) as cnt FROM properties WHERE city IS NOT NULL AND city != '' AND date(created_at) >= date(?) AND date(created_at) <= date(?) GROUP BY city ORDER BY cnt DESC LIMIT 8",
        params
    ).fetchall()

    # Deal pipeline by stage
    stage_order = ['offer', 'negotiation', 'under_contract', 'closed', 'lost']
    by_stage_raw = db.execute(
        "SELECT stage, COUNT(*) as cnt, COALESCE(SUM(value), 0) as total_value FROM deals WHERE date(created_at) >= date(?) AND date(created_at) <= date(?) GROUP BY stage",
        params
    ).fetchall()
    stage_map = {}
    for r in by_stage_raw:
        stage_map[r['stage']] = {
            'stage': r['stage'] or 'unknown',
            'cnt': int(r['cnt'] or 0),
            'total_value': float(r['total_value'] or 0),
        }
    by_stage = []
    for st in stage_order:
        if st in stage_map:
            by_stage.append(stage_map[st])
        else:
            by_stage.append({'stage': st, 'cnt': 0, 'total_value': 0.0})
    for st, row in stage_map.items():
        if st not in stage_order:
            by_stage.append(row)

    monthly = db.execute(
        "SELECT strftime('%Y-%m', created_at) as month, COUNT(*) as cnt FROM properties WHERE date(created_at) >= date(?) AND date(created_at) <= date(?) GROUP BY month ORDER BY month",
        params
    ).fetchall()
    monthly_deals = db.execute(
        "SELECT strftime('%Y-%m', updated_at) as month, COUNT(*) as cnt FROM deals WHERE stage = 'closed' AND date(updated_at) >= date(?) AND date(updated_at) <= date(?) GROUP BY month ORDER BY month",
        params
    ).fetchall()

    month_set = sorted({
        r['month'] for r in list(monthly) + list(monthly_deals)
        if r['month']
    })
    props_map = {r['month']: int(r['cnt'] or 0) for r in monthly if r['month']}
    deals_map = {r['month']: int(r['cnt'] or 0) for r in monthly_deals if r['month']}
    month_labels = list(month_set)
    month_props = [props_map.get(m, 0) for m in month_labels]
    month_deals = [deals_map.get(m, 0) for m in month_labels]

    type_labels = [(r['property_type'] or 'Other') for r in by_type]
    type_counts = [int(r['cnt'] or 0) for r in by_type]
    city_labels = [(r['city'] or '—') for r in by_city]
    city_counts = [int(r['cnt'] or 0) for r in by_city]
    stage_labels = [(r['stage'] or 'unknown').replace('_', ' ').title() for r in by_stage]
    stage_counts = [int(r['cnt'] or 0) for r in by_stage]
    stage_values = [float(r['total_value'] or 0) for r in by_stage]

    if export == 'csv':
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator='\r\n')

        # Title block
        w.writerow(['Estacle Real Estate CRM'])
        w.writerow(['Analytics Report'])
        w.writerow(['Report period', str(from_date), 'to', str(to_date)])
        w.writerow([])

        # Section A
        w.writerow(['SECTION A  KEY METRICS'])
        w.writerow(['Metric', 'Value', 'Description'])
        metrics = [
            ('Total Properties', total_properties, 'All properties in period'),
            ('Available Listings', available, 'Status = available'),
            ('Closed Deals', closed_deals, 'Deals closed in period'),
            ('Total Leads', total_leads, 'Clients and leads created'),
            ('Open Deals', open_deals, 'Active pipeline deals'),
            ('Pipeline Value (INR)', round(float(pipeline_value or 0), 2), 'Sum of open deal values'),
            ('Sold Properties', sold, 'Status = sold'),
            ('Rented Properties', rented, 'Status = rented'),
            ('Under Offer', under_offer, 'Status = under offer'),
            ('Appointments', total_appointments, 'Appointments created'),
            ('Agents', total_agents, 'Agent user accounts'),
        ]
        for m in metrics:
            w.writerow(list(m))
        w.writerow([])

        # Section B
        w.writerow(['SECTION B  DEAL PIPELINE'])
        w.writerow(['Stage', 'Deals', 'Value (INR)'])
        for r in by_stage:
            w.writerow([
                str(r.get('stage') or '').replace('_', ' ').title(),
                int(r.get('cnt') or 0),
                round(float(r.get('total_value') or 0), 2),
            ])
        w.writerow([])

        # Section C
        w.writerow(['SECTION C  PROPERTIES BY TYPE'])
        w.writerow(['Property Type', 'Count'])
        if type_labels:
            for lab, cnt in zip(type_labels, type_counts):
                w.writerow([lab, cnt])
        else:
            w.writerow(['No data', 0])
        w.writerow([])

        # Section D
        w.writerow(['SECTION D  TOP CITIES'])
        w.writerow(['City', 'Properties'])
        if city_labels:
            for lab, cnt in zip(city_labels, city_counts):
                w.writerow([lab, cnt])
        else:
            w.writerow(['No data', 0])
        w.writerow([])

        # Section E
        w.writerow(['SECTION E  MONTHLY PERFORMANCE'])
        w.writerow(['Month', 'Properties Added', 'Deals Closed'])
        if month_labels:
            for m, p, d in zip(month_labels, month_props, month_deals):
                w.writerow([m, p, d])
        else:
            w.writerow(['No data', 0, 0])
        w.writerow([])
        w.writerow(['End of report'])
        w.writerow(['Estacle CRM', 'Confidential  Internal use only'])

        content = '\ufeff' + buf.getvalue()
        resp = make_response(content.encode('utf-8'))
        resp.headers['Content-Type'] = 'text/csv; charset=utf-8'
        resp.headers['Content-Disposition'] = 'attachment; filename="estacle_reports.csv"'
        return resp



    max_type_count = max(type_counts) if type_counts else 1
    max_month_val = max(list(month_props) + list(month_deals)) if (month_props or month_deals) else 1
    if max_month_val < 1:
        max_month_val = 1

    if export == 'pdf':
        return render_template(
            'admin_analytics_report.html',
            stats=stats,
            by_type=by_type,
            by_city=by_city,
            by_stage=by_stage,
            from_date=from_date,
            to_date=to_date,
            type_labels=type_labels,
            type_counts=type_counts,
            stage_labels=stage_labels,
            stage_counts=stage_counts,
            stage_values=stage_values,
            month_labels=month_labels,
            month_props=month_props,
            month_deals=month_deals,
            max_type_count=max_type_count,
            max_month_val=max_month_val,
            format_money=format_money,
        )


    return render_template(
        'admin_analytics.html',
        stats=stats,
        by_type=by_type,
        by_city=by_city,
        by_stage=by_stage,
        from_date=from_date,
        to_date=to_date,
        type_labels=type_labels,
        type_counts=type_counts,
        city_labels=city_labels,
        city_counts=city_counts,
        stage_labels=stage_labels,
        stage_counts=stage_counts,
        stage_values=stage_values,
        month_labels=month_labels,
        month_props=month_props,
        month_deals=month_deals,
        format_money=format_money,
    )



@admin.route('/admin/reports')
@login_required
@role_required('admin')
def admin_reports():
    db = get_db()
    recent_props = db.execute('SELECT id, title, price, status, city, created_at FROM properties ORDER BY created_at DESC LIMIT 10').fetchall()
    recent_clients = db.execute('SELECT id, full_name, status, email, created_at FROM clients ORDER BY created_at DESC LIMIT 10').fetchall()
    agents = db.execute("""
        SELECT u.id, u.full_name,
               (SELECT COUNT(*) FROM properties WHERE agent_id = u.id) as prop_count,
               (SELECT COUNT(*) FROM clients WHERE agent_id = u.id) as client_count
        FROM users u WHERE u.role = 'agent' ORDER BY u.full_name
    """).fetchall()
    return render_template('admin_reports.html', recent_props=recent_props, recent_clients=recent_clients, agents=agents)


@admin.route('/admin/settings', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def admin_settings():
    if request.method == 'POST':
        currency = request.form.get('currency', 'INR')
        if currency not in ('USD', 'INR', 'AED'):
            flash('Invalid currency.', 'danger')
            return redirect(url_for('admin.admin_settings'))
        set_setting('currency', currency)
        set_setting('company_name', request.form.get('company_name', 'Estacle').strip() or 'Estacle')
        set_setting('support_email', request.form.get('support_email', '').strip())
        set_setting('support_phone', request.form.get('support_phone', '').strip())
        set_setting('timezone', request.form.get('timezone', 'Asia/Kolkata').strip() or 'Asia/Kolkata')
        set_setting('address', request.form.get('address', '').strip())
        flash('Settings saved.', 'success')
        return redirect(url_for('admin.admin_settings'))
    settings = {
        'currency': get_setting('currency', 'INR'),
        'company_name': get_setting('company_name', 'Estacle'),
        'support_email': get_setting('support_email', 'support@estacle.com'),
        'support_phone': get_setting('support_phone', ''),
        'timezone': get_setting('timezone', 'Asia/Kolkata'),
        'address': get_setting('address', ''),
    }
    return render_template('admin_settings.html', **settings)


# ---------- Property detail / edit ----------


@admin.route('/admin/properties/<int:prop_id>')
@login_required
@role_required('admin')
def admin_property_detail(prop_id):
    db = get_db()
    migrate_db()
    prop = db.execute(
        "SELECT p.*, u.full_name as agent_name FROM properties p LEFT JOIN users u ON p.agent_id = u.id WHERE p.id = ?",
        (prop_id,)
    ).fetchone()
    if not prop:
        flash('Property not found.', 'danger')
        return redirect(url_for('admin.admin_properties'))
    photos = []
    if prop['photos']:
        photos = [p.strip() for p in prop['photos'].split(',') if p.strip()]
    videos = []
    try:
        if prop['videos']:
            videos = [v.strip() for v in prop['videos'].split(',') if v.strip()]
    except (KeyError, IndexError, TypeError):
        videos = []
    currency = get_setting('currency', 'INR')
    return render_template(
        'admin_property_detail.html', prop=prop, photos=photos, videos=videos,
        currency=currency, format_money=format_money
    )


@admin.route('/admin/properties/edit/<int:prop_id>', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def admin_edit_property(prop_id):
    db = get_db()
    migrate_db()
    prop = db.execute('SELECT * FROM properties WHERE id = ?', (prop_id,)).fetchone()
    if not prop:
        flash('Property not found.', 'danger')
        return redirect(url_for('admin.admin_properties'))
    agents = db.execute(
        "SELECT id, full_name FROM users WHERE role = 'agent' AND is_active = 1 ORDER BY full_name"
    ).fetchall()
    currency = get_setting('currency', 'INR')

    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        if not title:
            flash('Title is required.', 'danger')
            return redirect(url_for('admin.admin_edit_property', prop_id=prop_id))
        try:
            price = float(request.form.get('price') or 0)
            bedrooms = int(request.form.get('bedrooms') or 0) or None
            bathrooms = float(request.form.get('bathrooms') or 0) or None
            area_sqft = float(request.form.get('area_sqft') or 0) or None
            carpet_area = float(request.form.get('carpet_area') or 0) or None
            monthly_rent = float(request.form.get('monthly_rent') or 0) or None
            security_deposit = float(request.form.get('security_deposit') or 0) or None
            maintenance = float(request.form.get('maintenance') or 0) or None
            per_sqft = float(request.form.get('per_sqft') or 0) or None
            agent_id = int(request.form.get('agent_id') or 0) or None
        except ValueError:
            flash('Invalid numeric values.', 'danger')
            return redirect(url_for('admin.admin_edit_property', prop_id=prop_id))

        photos_str = prop['photos'] if 'photos' in prop.keys() else None
        files = request.files.getlist('photos')
        photo_names = [p.strip() for p in (photos_str or '').split(',') if p.strip()]
        for f in files:
            if f and f.filename:
                fname = secure_filename(f.filename)
                unique = '{}_{}'.format(int(__import__('time').time()), fname)
                f.save(os.path.join(current_app.config['UPLOAD_FOLDER'], unique))
                photo_names.append(unique)
        if photo_names:
            photos_str = ','.join(photo_names)

        videos_str = prop['videos'] if 'videos' in prop.keys() else None
        vfiles = request.files.getlist('videos')
        video_names = [v.strip() for v in (videos_str or '').split(',') if v.strip()]
        for f in vfiles:
            if f and f.filename:
                fname = secure_filename(f.filename)
                ext = os.path.splitext(fname)[1].lower()
                if ext not in ('.mp4', '.webm', '.mov', '.ogg', '.m4v'):
                    continue
                unique = 'vid_{}_{}'.format(int(__import__('time').time()), fname)
                f.save(os.path.join(current_app.config['UPLOAD_FOLDER'], unique))
                video_names.append(unique)
        if video_names:
            videos_str = ','.join(video_names)

        db.execute(
            "UPDATE properties SET title=?, description=?, price=?, property_type=?, status=?, "
            "address=?, city=?, state=?, zip_code=?, bedrooms=?, bathrooms=?, area_sqft=?, agent_id=?, "
            "category=?, listing_type=?, floor=?, carpet_area=?, furnishing=?, parking=?, amenities=?, "
            "seller_name=?, seller_phone=?, monthly_rent=?, security_deposit=?, maintenance=?, per_sqft=?, "
            "latitude=?, longitude=?, photos=?, year_built=?, facing=?, videos=? WHERE id=?",
            (
                title,
                request.form.get('description', ''),
                price,
                request.form.get('property_type', 'house'),
                request.form.get('status', 'available'),
                request.form.get('address', ''),
                request.form.get('city', ''),
                request.form.get('state', ''),
                request.form.get('zip_code', ''),
                bedrooms, bathrooms, area_sqft, agent_id,
                request.form.get('category', 'sale'),
                request.form.get('listing_type', 'sale'),
                request.form.get('floor', '') or None,
                carpet_area,
                request.form.get('furnishing', '') or None,
                request.form.get('parking', '') or None,
                request.form.get('amenities', '') or None,
                request.form.get('seller_name', '') or None,
                request.form.get('seller_phone', '') or None,
                monthly_rent, security_deposit, maintenance, per_sqft,
                request.form.get('latitude', '') or None,
                request.form.get('longitude', '') or None,
                photos_str,
                request.form.get('year_built', '') or None,
                request.form.get('facing', '') or None,
                videos_str,
                prop_id,
            )
        )
        db.commit()
        log_activity('Updated property', 'property', prop_id, title)
        flash('Property updated.', 'success')
        return redirect(url_for('admin.admin_property_detail', prop_id=prop_id))

    return render_template(
        'admin_edit_property.html', prop=prop, agents=agents, currency=currency, format_money=format_money
    )


@admin.route('/admin/tasks')
@login_required
@role_required('admin')
def admin_tasks():
    db = get_db()
    tasks = db.execute(
        "SELECT t.*, u.full_name as assignee_name, c.full_name as client_name, p.title as property_title FROM tasks t LEFT JOIN users u ON t.assigned_to = u.id LEFT JOIN clients c ON t.related_client_id = c.id LEFT JOIN properties p ON t.related_property_id = p.id ORDER BY t.created_at DESC"
    ).fetchall()
    return render_template('admin_tasks.html', tasks=tasks)


@admin.route('/admin/tasks/add', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def admin_add_task():
    db = get_db()
    agents = db.execute("SELECT id, full_name FROM users WHERE role IN ('admin','agent') AND is_active=1 ORDER BY full_name").fetchall()
    clients = db.execute('SELECT id, full_name FROM clients ORDER BY full_name').fetchall()
    properties = db.execute('SELECT id, title FROM properties ORDER BY title').fetchall()
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        if not title:
            flash('Title required.', 'danger')
            return redirect(url_for('admin.admin_add_task'))
        assigned = request.form.get('assigned_to') or None
        client_id = request.form.get('related_client_id') or None
        prop_id = request.form.get('related_property_id') or None
        db.execute(
            "INSERT INTO tasks (title, description, due_date, status, priority, assigned_to, related_client_id, related_property_id, created_by) VALUES (?,?,?,?,?,?,?,?,?)",
            (title, request.form.get('description',''), request.form.get('due_date') or None,
             request.form.get('status','pending'), request.form.get('priority','medium'),
             int(assigned) if assigned else None, int(client_id) if client_id else None,
             int(prop_id) if prop_id else None, session['user_id'])
        )
        db.commit()
        log_activity('Created task', 'task', None, title)
        flash('Task created.', 'success')
        return redirect(url_for('admin.admin_tasks'))
    return render_template('admin_add_task.html', agents=agents, clients=clients, properties=properties)


@admin.route('/admin/tasks/status/<int:task_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_task_status(task_id):
    status = request.form.get('status', 'pending')
    if status not in ('pending', 'in_progress', 'done', 'cancelled'):
        flash('Invalid status.', 'danger')
        return redirect(url_for('admin.admin_tasks'))
    db = get_db()
    task = db.execute('SELECT title FROM tasks WHERE id=?', (task_id,)).fetchone()
    db.execute('UPDATE tasks SET status=? WHERE id=?', (status, task_id))
    db.commit()
    if task:
        log_activity('Updated task', 'task', task_id, f"{task['title']} → {status}")
    flash('Task updated.', 'success')
    if request.form.get('next') == 'dashboard':
        return redirect(url_for('admin.admin_dashboard'))
    return redirect(url_for('admin.admin_tasks'))


@admin.route('/admin/tasks/delete/<int:task_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_delete_task(task_id):
    db = get_db()
    db.execute('DELETE FROM tasks WHERE id=?', (task_id,))
    db.commit()
    flash('Task deleted.', 'success')
    return redirect(url_for('admin.admin_tasks'))


@admin.route('/admin/deals')
@login_required
@role_required('admin')
def admin_deals():
    db = get_db()
    deals = db.execute(
        "SELECT d.*, p.title as property_title, c.full_name as client_name, u.full_name as agent_name FROM deals d LEFT JOIN properties p ON d.property_id = p.id LEFT JOIN clients c ON d.client_id = c.id LEFT JOIN users u ON d.agent_id = u.id ORDER BY d.updated_at DESC"
    ).fetchall()
    return render_template('admin_deals.html', deals=deals)


@admin.route('/admin/deals/add', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def admin_add_deal():
    db = get_db()
    properties = db.execute('SELECT id, title FROM properties ORDER BY title').fetchall()
    clients = db.execute('SELECT id, full_name FROM clients ORDER BY full_name').fetchall()
    agents = db.execute("SELECT id, full_name FROM users WHERE role='agent' AND is_active=1 ORDER BY full_name").fetchall()
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        if not title:
            flash('Title required.', 'danger')
            return redirect(url_for('admin.admin_add_deal'))
        try:
            value = float(request.form.get('value') or 0)
            rate = float(request.form.get('commission_rate') or 2.5)
            prop_id = int(request.form.get('property_id') or 0) or None
            client_id = int(request.form.get('client_id') or 0) or None
            agent_id = int(request.form.get('agent_id') or 0) or None
        except ValueError:
            flash('Invalid numbers.', 'danger')
            return redirect(url_for('admin.admin_add_deal'))
        db.execute(
            "INSERT INTO deals (title, property_id, client_id, agent_id, stage, value, commission_rate, expected_close, notes) VALUES (?,?,?,?,?,?,?,?,?)",
            (title, prop_id, client_id, agent_id, request.form.get('stage','offer'), value, rate,
             request.form.get('expected_close') or None, request.form.get('notes',''))
        )
        db.commit()
        log_activity('Created deal', 'deal', None, title)
        flash('Deal created.', 'success')
        return redirect(url_for('admin.admin_deals'))
    return render_template('admin_add_deal.html', properties=properties, clients=clients, agents=agents)


@admin.route('/admin/deals/stage/<int:deal_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_deal_stage(deal_id):
    stage = request.form.get('stage', 'offer')
    if stage not in ('offer', 'negotiation', 'under_contract', 'closed', 'lost'):
        flash('Invalid stage.', 'danger')
        return redirect(url_for('admin.admin_deals'))
    db = get_db()
    deal = db.execute('SELECT title FROM deals WHERE id=?', (deal_id,)).fetchone()
    db.execute("UPDATE deals SET stage=?, updated_at=CURRENT_TIMESTAMP WHERE id=?", (stage, deal_id))
    db.commit()
    title = deal['title'] if deal else str(deal_id)
    log_activity('Changed deal stage', 'deal', deal_id, f'{title} → {stage}')
    flash('Deal stage updated.', 'success')
    return redirect(url_for('admin.admin_deals'))


@admin.route('/admin/deals/delete/<int:deal_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_delete_deal(deal_id):
    db = get_db()
    db.execute('DELETE FROM deals WHERE id=?', (deal_id,))
    db.commit()
    flash('Deal deleted.', 'success')
    return redirect(url_for('admin.admin_deals'))


@admin.route('/admin/activity')
@login_required
@role_required('admin')
def admin_activity():
    db = get_db()
    logs = db.execute(
        "SELECT a.*, u.full_name as user_name FROM activity_log a LEFT JOIN users u ON a.user_id = u.id ORDER BY a.created_at DESC LIMIT 100"
    ).fetchall()
    return render_template('admin_activity.html', logs=logs)


@admin.route('/admin/documents')
@login_required
@role_required('admin')
def admin_documents():
    db = get_db()
    docs = db.execute(
        "SELECT d.*, u.full_name as creator_name FROM documents d LEFT JOIN users u ON d.created_by = u.id ORDER BY d.created_at DESC"
    ).fetchall()
    return render_template('admin_documents.html', docs=docs)


@admin.route('/admin/documents/add', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def admin_add_document():
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        if not title:
            flash('Title required.', 'danger')
            return redirect(url_for('admin.admin_add_document'))
        db = get_db()
        rid = request.form.get('related_id') or None
        db.execute(
            "INSERT INTO documents (title, doc_type, related_type, related_id, expiry_date, notes, created_by) VALUES (?,?,?,?,?,?,?)",
            (title, request.form.get('doc_type',''), request.form.get('related_type',''),
             int(rid) if rid else None, request.form.get('expiry_date') or None,
             request.form.get('notes',''), session['user_id'])
        )
        db.commit()
        log_activity('Added document', 'document', None, title)
        flash('Document recorded.', 'success')
        return redirect(url_for('admin.admin_documents'))
    return render_template('admin_add_document.html')


@admin.route('/admin/documents/delete/<int:doc_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_delete_document(doc_id):
    db = get_db()
    db.execute('DELETE FROM documents WHERE id=?', (doc_id,))
    db.commit()
    flash('Document removed.', 'success')
    return redirect(url_for('admin.admin_documents'))


@admin.route('/admin/properties/display')
@login_required
@role_required('admin')
def admin_display_properties():
    db = get_db()
    migrate_db()
    q = request.args.get('q', '').strip()
    listing_type = request.args.get('listing_type', '').strip()
    prop_type = request.args.get('property_type', '').strip()
    city = request.args.get('city', '').strip()
    status = request.args.get('status', '').strip()
    beds = request.args.get('beds', '').strip()
    pmin = request.args.get('pmin', '').strip()
    pmax = request.args.get('pmax', '').strip()

    sql = "SELECT p.*, u.full_name as agent_name FROM properties p LEFT JOIN users u ON p.agent_id = u.id WHERE 1=1"
    params = []
    if q:
        sql += " AND (p.title LIKE ? OR p.city LIKE ? OR p.address LIKE ? OR p.description LIKE ?)"
        params.extend([f'%{q}%'] * 4)
    if listing_type in ('sale', 'rent'):
        sql += " AND p.listing_type = ?"
        params.append(listing_type)
    if prop_type:
        sql += " AND p.property_type = ?"
        params.append(prop_type)
    if city:
        sql += " AND p.city LIKE ?"
        params.append(f'%{city}%')
    if status:
        sql += " AND p.status = ?"
        params.append(status)
    if beds:
        try:
            sql += " AND p.bedrooms >= ?"
            params.append(int(beds))
        except ValueError:
            pass
    if pmin:
        try:
            sql += " AND p.price >= ?"
            params.append(float(pmin))
        except ValueError:
            pass
    if pmax:
        try:
            sql += " AND p.price <= ?"
            params.append(float(pmax))
        except ValueError:
            pass
    sql += " ORDER BY p.created_at DESC"
    properties = db.execute(sql, params).fetchall()
    currency = get_setting('currency', 'INR')
    return render_template(
        'admin_display_properties.html',
        properties=properties,
        currency=currency,
        format_money=format_money,
        filters={'q': q, 'listing_type': listing_type, 'property_type': prop_type, 'city': city, 'status': status, 'beds': beds, 'pmin': pmin, 'pmax': pmax}
    )


@admin.route('/admin/search')
@login_required
@role_required('admin')
def admin_search():
    q = request.args.get('q', '').strip()
    stype = request.args.get('type', '').strip().lower()
    valid_types = {
        'users', 'properties', 'clients', 'deals', 'leads',
        'inquiries', 'appointments', 'tasks', 'documents', 'activity',
    }
    if stype not in valid_types:
        stype = ''

    db = get_db()
    migrate_db()
    results = {
        'users': [], 'properties': [], 'clients': [], 'deals': [],
        'leads': [], 'inquiries': [], 'appointments': [], 'tasks': [],
        'documents': [], 'activity': [],
    }
    counts = {
        'users': db.execute('SELECT COUNT(*) FROM users').fetchone()[0],
        'properties': db.execute('SELECT COUNT(*) FROM properties').fetchone()[0],
        'clients': db.execute('SELECT COUNT(*) FROM clients').fetchone()[0],
        'deals': db.execute('SELECT COUNT(*) FROM deals').fetchone()[0],
        'leads': db.execute('SELECT COUNT(*) FROM leads').fetchone()[0],
        'inquiries': db.execute('SELECT COUNT(*) FROM inquiries').fetchone()[0],
        'appointments': db.execute('SELECT COUNT(*) FROM appointments').fetchone()[0],
        'tasks': db.execute('SELECT COUNT(*) FROM tasks').fetchone()[0],
        'documents': db.execute('SELECT COUNT(*) FROM documents').fetchone()[0],
        'activity': db.execute('SELECT COUNT(*) FROM activity_log').fetchone()[0],
    }

    like = f'%{q}%' if q else None
    num = None
    if q:
        q_clean = q.replace(',', '').replace('₹', '').replace('$', '').replace('AED', '').strip()
        try:
            num = float(q_clean)
        except ValueError:
            num = None

    def want(key):
        return (not stype) or stype == key

    # Load data when searching or when a category type is selected
    if q or stype:
        if want('users'):
            if q:
                results['users'] = db.execute(
                    "SELECT id, full_name, email, role, phone FROM users "
                    "WHERE full_name LIKE ? OR email LIKE ? OR phone LIKE ? LIMIT 50",
                    (like, like, like)
                ).fetchall()
            else:
                results['users'] = db.execute(
                    "SELECT id, full_name, email, role, phone FROM users ORDER BY created_at DESC LIMIT 50"
                ).fetchall()

        if want('properties'):
            if q:
                prop_sql = (
                    "SELECT id, title, city, price, status FROM properties "
                    "WHERE title LIKE ? OR city LIKE ? OR address LIKE ? OR seller_name LIKE ?"
                )
                prop_params = [like, like, like, like]
                if num is not None:
                    prop_sql += (
                        " OR price = ? OR CAST(price AS TEXT) LIKE ? OR bedrooms = ? "
                        "OR bathrooms = ? OR area_sqft = ? OR monthly_rent = ?"
                    )
                    prop_params.extend([num, like, int(num) if num == int(num) else -1, num, num, num])
                prop_sql += " LIMIT 50"
                results['properties'] = db.execute(prop_sql, prop_params).fetchall()
            else:
                results['properties'] = db.execute(
                    "SELECT id, title, city, price, status FROM properties ORDER BY created_at DESC LIMIT 50"
                ).fetchall()

        if want('clients'):
            if q:
                results['clients'] = db.execute(
                    "SELECT id, full_name, email, phone, status FROM clients "
                    "WHERE full_name LIKE ? OR email LIKE ? OR phone LIKE ? LIMIT 50",
                    (like, like, like)
                ).fetchall()
            else:
                results['clients'] = db.execute(
                    "SELECT id, full_name, email, phone, status FROM clients ORDER BY created_at DESC LIMIT 50"
                ).fetchall()

        if want('deals'):
            if q:
                deal_sql = "SELECT id, title, stage, value FROM deals WHERE title LIKE ? OR notes LIKE ?"
                deal_params = [like, like]
                if num is not None:
                    deal_sql += " OR value = ? OR CAST(value AS TEXT) LIKE ? OR commission_rate = ?"
                    deal_params.extend([num, like, num])
                deal_sql += " LIMIT 50"
                results['deals'] = db.execute(deal_sql, deal_params).fetchall()
            else:
                results['deals'] = db.execute(
                    "SELECT id, title, stage, value FROM deals ORDER BY created_at DESC LIMIT 50"
                ).fetchall()

        if want('leads'):
            if q:
                lead_sql = (
                    "SELECT id, full_name, email, phone, source, interest, status, budget FROM leads "
                    "WHERE full_name LIKE ? OR email LIKE ? OR phone LIKE ? OR source LIKE ? "
                    "OR interest LIKE ? OR notes LIKE ?"
                )
                lead_params = [like, like, like, like, like, like]
                if num is not None:
                    lead_sql += " OR budget = ? OR CAST(budget AS TEXT) LIKE ?"
                    lead_params.extend([num, like])
                lead_sql += " LIMIT 50"
                results['leads'] = db.execute(lead_sql, lead_params).fetchall()
            else:
                results['leads'] = db.execute(
                    "SELECT id, full_name, email, phone, source, interest, status, budget FROM leads "
                    "ORDER BY created_at DESC LIMIT 50"
                ).fetchall()

        if want('inquiries'):
            if q:
                results['inquiries'] = db.execute(
                    "SELECT i.id, i.full_name, i.email, i.phone, i.message, i.status, i.property_id, "
                    "p.title AS property_title FROM inquiries i "
                    "LEFT JOIN properties p ON i.property_id = p.id "
                    "WHERE i.full_name LIKE ? OR i.email LIKE ? OR i.phone LIKE ? OR i.message LIKE ? LIMIT 50",
                    (like, like, like, like)
                ).fetchall()
            else:
                results['inquiries'] = db.execute(
                    "SELECT i.id, i.full_name, i.email, i.phone, i.message, i.status, i.property_id, "
                    "p.title AS property_title FROM inquiries i "
                    "LEFT JOIN properties p ON i.property_id = p.id ORDER BY i.created_at DESC LIMIT 50"
                ).fetchall()

        if want('appointments'):
            if q:
                results['appointments'] = db.execute(
                    "SELECT a.id, a.appointment_date, a.status, a.notes, "
                    "c.full_name AS client_name, p.title AS property_title "
                    "FROM appointments a "
                    "LEFT JOIN clients c ON a.client_id = c.id "
                    "LEFT JOIN properties p ON a.property_id = p.id "
                    "WHERE a.notes LIKE ? OR c.full_name LIKE ? OR p.title LIKE ? LIMIT 50",
                    (like, like, like)
                ).fetchall()
            else:
                results['appointments'] = db.execute(
                    "SELECT a.id, a.appointment_date, a.status, a.notes, "
                    "c.full_name AS client_name, p.title AS property_title "
                    "FROM appointments a "
                    "LEFT JOIN clients c ON a.client_id = c.id "
                    "LEFT JOIN properties p ON a.property_id = p.id "
                    "ORDER BY a.appointment_date DESC LIMIT 50"
                ).fetchall()

        if want('tasks'):
            if q:
                results['tasks'] = db.execute(
                    "SELECT id, title, description, due_date, status, priority FROM tasks "
                    "WHERE title LIKE ? OR description LIKE ? LIMIT 50",
                    (like, like)
                ).fetchall()
            else:
                results['tasks'] = db.execute(
                    "SELECT id, title, description, due_date, status, priority FROM tasks "
                    "ORDER BY created_at DESC LIMIT 50"
                ).fetchall()

        if want('documents'):
            if q:
                results['documents'] = db.execute(
                    "SELECT id, title, doc_type, related_type, expiry_date, notes FROM documents "
                    "WHERE title LIKE ? OR doc_type LIKE ? OR notes LIKE ? LIMIT 50",
                    (like, like, like)
                ).fetchall()
            else:
                results['documents'] = db.execute(
                    "SELECT id, title, doc_type, related_type, expiry_date, notes FROM documents "
                    "ORDER BY created_at DESC LIMIT 50"
                ).fetchall()

        if want('activity'):
            if q:
                results['activity'] = db.execute(
                    "SELECT a.id, a.action, a.entity_type, a.entity_id, a.details, a.created_at, "
                    "u.full_name AS user_name FROM activity_log a "
                    "LEFT JOIN users u ON a.user_id = u.id "
                    "WHERE a.action LIKE ? OR a.details LIKE ? OR a.entity_type LIKE ? "
                    "OR u.full_name LIKE ? LIMIT 50",
                    (like, like, like, like)
                ).fetchall()
            else:
                results['activity'] = db.execute(
                    "SELECT a.id, a.action, a.entity_type, a.entity_id, a.details, a.created_at, "
                    "u.full_name AS user_name FROM activity_log a "
                    "LEFT JOIN users u ON a.user_id = u.id ORDER BY a.created_at DESC LIMIT 50"
                ).fetchall()

    currency = get_setting('currency', 'INR')
    total = sum(len(v) for v in results.values())
    return render_template(
        'admin_search.html', q=q, stype=stype, results=results, total=total,
        counts=counts, currency=currency, format_money=format_money
    )


@admin.route('/admin/leads')
@login_required
@role_required('admin')
def admin_leads():
    db = get_db()
    migrate_db()
    leads = db.execute(
        "SELECT l.*, u.full_name as agent_name FROM leads l LEFT JOIN users u ON l.assigned_to = u.id ORDER BY l.created_at DESC"
    ).fetchall()
    return render_template('admin_leads.html', leads=leads)


@admin.route('/admin/leads/add', methods=['GET', 'POST'])
@login_required
@role_required('admin')
def admin_add_lead():
    db = get_db()
    migrate_db()
    agents = db.execute("SELECT id, full_name FROM users WHERE role IN ('admin','agent') AND is_active=1 ORDER BY full_name").fetchall()
    if request.method == 'POST':
        name = request.form.get('full_name', '').strip()
        if not name:
            flash('Name required.', 'danger')
            return redirect(url_for('admin.admin_add_lead'))
        assigned = request.form.get('assigned_to') or None
        try:
            budget = float(request.form.get('budget') or 0) or None
        except ValueError:
            budget = None
        db.execute(
            "INSERT INTO leads (full_name, email, phone, source, interest, budget, status, notes, assigned_to) VALUES (?,?,?,?,?,?,?,?,?)",
            (name, request.form.get('email',''), request.form.get('phone',''), request.form.get('source',''),
             request.form.get('interest',''), budget, request.form.get('status','new'), request.form.get('notes',''),
             int(assigned) if assigned else None)
        )
        db.commit()
        log_activity('Created lead', 'lead', None, name)
        flash('Lead added.', 'success')
        return redirect(url_for('admin.admin_leads'))
    return render_template('admin_add_lead.html', agents=agents)


@admin.route('/admin/leads/status/<int:lead_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_lead_status(lead_id):
    status = request.form.get('status', 'new')
    if status not in ('new', 'contacted', 'qualified', 'converted', 'lost'):
        flash('Invalid status.', 'danger')
        return redirect(url_for('admin.admin_leads'))
    db = get_db()
    db.execute('UPDATE leads SET status=? WHERE id=?', (status, lead_id))
    db.commit()
    flash('Lead updated.', 'success')
    return redirect(url_for('admin.admin_leads'))


@admin.route('/admin/leads/delete/<int:lead_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_delete_lead(lead_id):
    db = get_db()
    db.execute('DELETE FROM leads WHERE id=?', (lead_id,))
    db.commit()
    flash('Lead deleted.', 'success')
    return redirect(url_for('admin.admin_leads'))


@admin.route('/admin/commission')
@login_required
@role_required('admin')
def admin_commission():
    db = get_db()
    deals = db.execute(
        """SELECT d.*, p.title as property_title, u.full_name as agent_name
           FROM deals d LEFT JOIN properties p ON d.property_id = p.id
           LEFT JOIN users u ON d.agent_id = u.id
           WHERE d.stage = 'closed'
           ORDER BY d.updated_at DESC"""
    ).fetchall()
    total_value = sum((d['value'] or 0) for d in deals)
    total_comm = sum(((d['value'] or 0) * (d['commission_rate'] or 0) / 100) for d in deals)
    currency = get_setting('currency', 'INR')
    return render_template(
        'admin_commission.html', deals=deals, total_value=total_value, total_comm=total_comm,
        currency=currency, format_money=format_money
    )


@admin.route('/admin/compare')
@login_required
@role_required('admin')
def admin_compare():
    db = get_db()
    migrate_db()
    all_props = db.execute("SELECT id, title FROM properties ORDER BY title").fetchall()
    ids = request.args.getlist('ids')
    props = []
    if ids:
        try:
            id_list = [int(i) for i in ids][:4]
            if id_list:
                placeholders = ','.join('?' * len(id_list))
                props = db.execute(
                    f"SELECT * FROM properties WHERE id IN ({placeholders})", id_list
                ).fetchall()
        except ValueError:
            pass
    currency = get_setting('currency', 'INR')
    return render_template('admin_compare.html', all_props=all_props, props=props, currency=currency, format_money=format_money)


@admin.route('/admin/properties/bulk', methods=['POST'])
@login_required
@role_required('admin')
def admin_bulk_properties():
    db = get_db()
    action = request.form.get('action', '')
    ids = request.form.getlist('ids')
    if not ids:
        flash('No properties selected.', 'warning')
        return redirect(url_for('admin.admin_properties'))
    try:
        ids = [int(i) for i in ids]
    except ValueError:
        flash('Invalid selection.', 'danger')
        return redirect(url_for('admin.admin_properties'))
    placeholders = ','.join('?' * len(ids))
    if action == 'delete':
        db.execute(f'DELETE FROM properties WHERE id IN ({placeholders})', ids)
        db.commit()
        log_activity('Bulk deleted properties', 'property', None, f'{len(ids)} items')
        flash(f'Deleted {len(ids)} properties.', 'success')
    elif action in ('available', 'under_offer', 'sold', 'rented'):
        db.execute(f"UPDATE properties SET status = ? WHERE id IN ({placeholders})", [action] + ids)
        db.commit()
        log_activity('Bulk status update', 'property', None, f'{len(ids)} → {action}')
        flash(f'Updated status of {len(ids)} properties to {action}.', 'success')
    else:
        flash('Unknown action.', 'warning')
    return redirect(url_for('admin.admin_properties'))


@admin.route('/admin/export/properties')
@login_required
@role_required('admin')
def admin_export_properties():
    import csv
    import io
    from flask import Response
    db = get_db()
    rows = db.execute('''
        SELECT p.id, p.title, p.property_type, p.listing_type, p.status, p.price, p.city, p.state,
               p.bedrooms, p.bathrooms, p.area_sqft, u.full_name as agent_name, p.created_at
        FROM properties p LEFT JOIN users u ON p.agent_id = u.id
        ORDER BY p.id
    ''').fetchall()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['ID', 'Title', 'Type', 'Listing', 'Status', 'Price', 'City', 'State',
                     'Bedrooms', 'Bathrooms', 'Area Sqft', 'Agent', 'Created'])
    for r in rows:
        writer.writerow([r['id'], r['title'], r['property_type'], r['listing_type'], r['status'],
                         r['price'], r['city'], r['state'], r['bedrooms'], r['bathrooms'],
                         r['area_sqft'], r['agent_name'], r['created_at']])
    output.seek(0)
    return Response(
        output.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': 'attachment; filename=estacle_properties.csv'}
    )


@admin.route('/admin/export/clients')
@login_required
@role_required('admin')
def admin_export_clients():
    import csv
    import io
    from flask import Response
    db = get_db()
    rows = db.execute('''
        SELECT c.id, c.full_name, c.email, c.phone, c.status, u.full_name as agent_name, c.created_at
        FROM clients c LEFT JOIN users u ON c.agent_id = u.id ORDER BY c.id
    ''').fetchall()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['ID', 'Full Name', 'Email', 'Phone', 'Status', 'Agent', 'Created'])
    for r in rows:
        writer.writerow([r['id'], r['full_name'], r['email'], r['phone'], r['status'], r['agent_name'], r['created_at']])
    output.seek(0)
    return Response(
        output.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': 'attachment; filename=estacle_clients.csv'}
    )


@admin.route('/admin/export/users')
@login_required
@role_required('admin')
def admin_export_users():
    import csv
    import io
    from flask import Response
    db = get_db()
    rows = db.execute(
        'SELECT id, full_name, email, role, phone, is_active, created_at FROM users ORDER BY id'
    ).fetchall()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['ID', 'Full Name', 'Email', 'Role', 'Phone', 'Active', 'Created'])
    for r in rows:
        writer.writerow([r['id'], r['full_name'], r['email'], r['role'], r['phone'], r['is_active'], r['created_at']])
    output.seek(0)
    return Response(
        output.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': 'attachment; filename=estacle_users.csv'}
    )


@admin.route('/admin/export/deals')
@login_required
@role_required('admin')
def admin_export_deals():
    import csv
    import io
    from flask import Response
    db = get_db()
    rows = db.execute('''
        SELECT d.id, d.title, d.stage, d.value, d.commission_rate, d.expected_close,
               p.title as property_title, c.full_name as client_name, u.full_name as agent_name, d.created_at
        FROM deals d
        LEFT JOIN properties p ON d.property_id = p.id
        LEFT JOIN clients c ON d.client_id = c.id
        LEFT JOIN users u ON d.agent_id = u.id
        ORDER BY d.id
    ''').fetchall()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['ID', 'Title', 'Stage', 'Value', 'Commission %', 'Expected Close',
                     'Property', 'Client', 'Agent', 'Created'])
    for r in rows:
        writer.writerow([r['id'], r['title'], r['stage'], r['value'], r['commission_rate'],
                         r['expected_close'], r['property_title'], r['client_name'],
                         r['agent_name'], r['created_at']])
    output.seek(0)
    return Response(
        output.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': 'attachment; filename=estacle_deals.csv'}
    )


@admin.route('/admin/inquiries')
@login_required
@role_required('admin')
def admin_inquiries():
    db = get_db()
    migrate_db()
    rows = db.execute('''
        SELECT i.*, p.title as property_title
        FROM inquiries i
        LEFT JOIN properties p ON i.property_id = p.id
        ORDER BY i.created_at DESC
    ''').fetchall()
    return render_template('admin_inquiries.html', inquiries=rows)


@admin.route('/admin/inquiries/status/<int:inq_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_inquiry_status(inq_id):
    status = request.form.get('status', 'new')
    if status not in ('new', 'contacted', 'closed'):
        flash('Invalid status.', 'danger')
        return redirect(url_for('admin.admin_inquiries'))
    db = get_db()
    db.execute('UPDATE inquiries SET status = ? WHERE id = ?', (status, inq_id))
    db.commit()
    flash('Inquiry status updated.', 'success')
    return redirect(url_for('admin.admin_inquiries'))


@admin.route('/admin/inquiries/delete/<int:inq_id>', methods=['POST'])
@login_required
@role_required('admin')
def admin_delete_inquiry(inq_id):
    db = get_db()
    db.execute('DELETE FROM inquiries WHERE id = ?', (inq_id,))
    db.commit()
    flash('Inquiry deleted.', 'info')
    return redirect(url_for('admin.admin_inquiries'))



